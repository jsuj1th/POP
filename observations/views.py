import io
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from django.views.decorators.http import require_GET
import openpyxl

from .models import School, Teacher, Observation, ObservationSession
from .forms import ObservationForm


def login_view(request):
    if request.user.is_authenticated:
        return redirect('select_data')

    error = None
    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        password = request.POST.get('password', '')
        user = authenticate(request, username=username, password=password)
        if user is not None:
            login(request, user)
            return redirect('select_data')
        else:
            error = 'Invalid username or password.'

    return render(request, 'observations/login.html', {'error': error})


def logout_view(request):
    logout(request)
    return redirect('login')


@login_required
def select_data(request):
    if request.method == 'POST':
        school_id = request.POST.get('school_id', '').strip()
        teacher_id = request.POST.get('teacher_id', '').strip()

        if not school_id:
            messages.error(request, 'Please select a school.')
        elif not teacher_id:
            messages.error(request, 'Please select a teacher.')
        else:
            try:
                school = School.objects.get(pk=school_id)
                teacher = Teacher.objects.get(pk=teacher_id, school=school)
                return redirect(f'/observe/?school_id={school_id}&teacher_id={teacher_id}')
            except (School.DoesNotExist, Teacher.DoesNotExist):
                messages.error(request, 'Invalid school or teacher selection.')

    schools = School.objects.all().order_by('name')
    return render(request, 'observations/select_data.html', {'schools': schools})


@require_GET
def get_teachers(request):
    school_id = request.GET.get('school_id', '')
    if not school_id:
        return JsonResponse({'teachers': []})
    try:
        school = School.objects.get(pk=school_id)
        teachers = list(
            school.teachers.order_by('name').values('id', 'name', 'teacher_id')
        )
        return JsonResponse({'teachers': teachers})
    except School.DoesNotExist:
        return JsonResponse({'teachers': []})


@login_required
def observation_form(request):
    school_id = request.GET.get('school_id') or request.POST.get('school_id')
    teacher_id = request.GET.get('teacher_id') or request.POST.get('teacher_id')
    session_id = request.GET.get('session_id') or request.POST.get('session_id')
    from_reports = request.GET.get('from') == 'reports' or request.POST.get('from_reports') == '1'

    if not school_id or not teacher_id:
        messages.error(request, 'Please select a school and teacher first.')
        return redirect('select_data')

    school = get_object_or_404(School, pk=school_id)
    teacher = get_object_or_404(Teacher, pk=teacher_id, school=school)

    if session_id and not from_reports:
        # session_id in URL from a mid-session redirect — only reuse if still open
        session = get_object_or_404(ObservationSession, pk=session_id, observer=request.user, teacher=teacher)
        if session.completed:
            # Session was completed; fall through to find/create a fresh one
            session = None
    elif session_id and from_reports:
        # Coming from reports — target a specific session regardless of status
        session = get_object_or_404(ObservationSession, pk=session_id, observer=request.user, teacher=teacher)
    else:
        session = None

    if session is None:
        # Resume the most recent uncompleted session, or start a new one
        session = ObservationSession.objects.filter(
            observer=request.user,
            teacher=teacher,
            completed=False,
        ).order_by('-created_at').first()

        if not session:
            session = ObservationSession.objects.create(
                observer=request.user,
                school=school,
                teacher=teacher,
            )

    existing_count = Observation.objects.filter(session=session).count()

    if existing_count >= 60:
        session.completed = True
        session.save()
        messages.warning(request, 'Maximum of 60 observations reached. Starting a new session.')
        return redirect(f'/observe/?school_id={school_id}&teacher_id={teacher_id}')

    observation_number = existing_count + 1

    if request.method == 'POST':
        action = request.POST.get('action', 'enter')

        if action == 'complete':
            session.completed = True
            session.save()
            return redirect(f'/observe/complete/?session_id={session.id}')

        form = ObservationForm(request.POST)

        if form.is_valid():
            existing_count = Observation.objects.filter(session=session).count()

            if existing_count >= 60:
                messages.error(request, 'You cannot add more than 60 observations.')
                return redirect(f'/observe/complete/?session_id={session.id}')

            obs = form.save(commit=False)
            obs.observer = request.user
            obs.school = school
            obs.teacher = teacher
            obs.session = session
            obs.observation_number = existing_count + 1
            obs.save()

            if from_reports:
                messages.success(request, f'Observation #{obs.observation_number} for {teacher.name} added.')
                return redirect('reports')

            return redirect(f'/observe/?school_id={school_id}&teacher_id={teacher_id}&session_id={session.id}')

    else:
        form = ObservationForm()

    return render(request, 'observations/observation_form.html', {
        'form': form,
        'school': school,
        'teacher': teacher,
        'observation_number': observation_number,
        'school_id': school_id,
        'teacher_id': teacher_id,
        'session_id': session.id,
        'from_reports': from_reports,
    })


@login_required
def observation_complete(request):
    session_id = request.GET.get('session_id')
    session = get_object_or_404(ObservationSession, pk=session_id)

    observations = Observation.objects.filter(
        session=session
    ).order_by('observation_number')

    return render(request, 'observations/observation_complete.html', {
        'teacher': session.teacher,
        'observations': observations,
        'count': observations.count(),
        'session_id': session.id
    })


@login_required
def edit_observation(request, obs_id):
    obs = get_object_or_404(Observation, pk=obs_id, observer=request.user)
    if request.method == 'POST':
        form = ObservationForm(request.POST, instance=obs)
        if form.is_valid():
            form.save()
            messages.success(request, f'Observation #{obs.observation_number} updated.')
            return redirect('reports')
    else:
        form = ObservationForm(instance=obs)

    return render(request, 'observations/edit_observation.html', {
        'form': form,
        'obs': obs
    })


@login_required
def reports(request):
    observations = Observation.objects.filter(observer=request.user).select_related(
        'school', 'teacher', 'session'
    ).order_by('session__id', 'observation_number')

    groups = {}
    for obs in observations:
        sid = obs.session.id
        if sid not in groups:
            groups[sid] = {
                'session': obs.session,
                'teacher': obs.teacher,
                'school': obs.school,
                'observations': [],
            }
        groups[sid]['observations'].append(obs)

    return render(request, 'observations/reports.html', {'groups': groups.values()})

@login_required
def download_excel(request):
    observations = Observation.objects.filter(observer=request.user).select_related(
        'school', 'teacher', 'session'
    ).order_by('-created_at')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Observations'

    headers = [
        'Session', 'School', 'Teacher', 'Teacher ID', 'Obs #',
        'ESL Strategy', 'Curriculum', 'Physical Group',
        'Activity Structure', 'Mode', 'Language Content',
        'Lang Instr (T)', 'Lang Instr (S)', 'Date'
    ]

    for col_num, header in enumerate(headers, 1):
        ws.cell(row=1, column=col_num, value=header)

    for row_num, obs in enumerate(observations, 2):
        ws.cell(row=row_num, column=1, value=obs.session.id)
        ws.cell(row=row_num, column=2, value=obs.school.name)
        ws.cell(row=row_num, column=3, value=obs.teacher.name)
        ws.cell(row=row_num, column=4, value=obs.teacher.teacher_id)
        ws.cell(row=row_num, column=5, value=obs.observation_number)
        ws.cell(row=row_num, column=6, value=obs.esl_strategy)
        ws.cell(row=row_num, column=7, value=obs.curriculum)
        ws.cell(row=row_num, column=8, value=obs.physical_group)
        ws.cell(row=row_num, column=9, value=obs.activity_structure)
        ws.cell(row=row_num, column=10, value=obs.mode)
        ws.cell(row=row_num, column=11, value=obs.language_content)
        ws.cell(row=row_num, column=12, value=obs.language_instruction_teacher)
        ws.cell(row=row_num, column=13, value=obs.language_instruction_student)
        ws.cell(row=row_num, column=14, value=obs.created_at.strftime('%Y-%m-%d %H:%M'))

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="observations_all_sessions.xlsx"'
    return response

@login_required
def download_teacher_excel(request):
    session_id = request.GET.get('session_id')
    session = get_object_or_404(ObservationSession, pk=session_id)

    observations = Observation.objects.filter(
        session=session
    ).select_related('school', 'teacher').order_by('observation_number')

    wb = openpyxl.Workbook()
    ws = wb.active

    headers = [
        'School', 'Teacher', 'Teacher ID', 'Obs #',
        'ESL Strategy', 'Curriculum', 'Physical Group',
        'Activity Structure', 'Mode', 'Language Content',
        'Lang Instr (T)', 'Lang Instr (S)', 'Date'
    ]

    for col_num, header in enumerate(headers, 1):
        ws.cell(row=1, column=col_num, value=header)

    for row_num, obs in enumerate(observations, 2):
        ws.cell(row=row_num, column=1, value=obs.school.name)
        ws.cell(row=row_num, column=2, value=obs.teacher.name)
        ws.cell(row=row_num, column=3, value=obs.teacher.teacher_id)
        ws.cell(row=row_num, column=4, value=obs.observation_number)
        ws.cell(row=row_num, column=5, value=obs.esl_strategy)
        ws.cell(row=row_num, column=6, value=obs.curriculum)
        ws.cell(row=row_num, column=7, value=obs.physical_group)
        ws.cell(row=row_num, column=8, value=obs.activity_structure)
        ws.cell(row=row_num, column=9, value=obs.mode)
        ws.cell(row=row_num, column=10, value=obs.language_content)
        ws.cell(row=row_num, column=11, value=obs.language_instruction_teacher)
        ws.cell(row=row_num, column=12, value=obs.language_instruction_student)
        ws.cell(row=row_num, column=13, value=obs.created_at.strftime('%Y-%m-%d %H:%M'))

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="session_{session.id}.xlsx"'
    return response

@login_required
def select_session_download(request):
    teacher_id = request.GET.get('teacher_id')

    teacher = get_object_or_404(Teacher, pk=teacher_id)

    sessions = ObservationSession.objects.filter(
        observer=request.user,
        teacher=teacher
    ).order_by('-created_at')

    return render(request, 'observations/select_session_download.html', {
        'teacher': teacher,
        'sessions': sessions
    })
